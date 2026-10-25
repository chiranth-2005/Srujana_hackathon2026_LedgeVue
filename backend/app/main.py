import csv
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import uuid
from collections.abc import Iterable, Iterator
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from itertools import chain, islice
from math import ceil
from typing import Annotated, BinaryIO, Literal
from zipfile import BadZipFile

import xlrd
from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile, status
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, EmailStr, Field, StringConstraints, field_validator
from pwdlib import PasswordHash
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address
from starlette.concurrency import run_in_threadpool
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, JSON, Numeric, String, Text, UniqueConstraint, create_engine, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, sessionmaker


APP_ENV = os.getenv("APP_ENV", "development").lower()
DATABASE_URL = os.getenv("DATABASE_URL") or "sqlite:///./ledgevue_local.db"
if DATABASE_URL.startswith("postgresql+psycopg://"):
    engine = create_engine(DATABASE_URL, pool_pre_ping=True, pool_recycle=300)
elif DATABASE_URL.startswith("sqlite"):
    engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
else:
    raise RuntimeError("DATABASE_URL must use PostgreSQL with the psycopg driver or SQLite")

COOKIE_SECURE = os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true"
if APP_ENV == "production" and not COOKIE_SECURE:
    raise RuntimeError("SESSION_COOKIE_SECURE=true is required in production")

APP_ORIGINS = {value.strip().rstrip("/") for value in os.getenv("APP_ORIGINS", "").split(",") if value.strip()}
ALLOWED_HOSTS = [value.strip() for value in os.getenv("ALLOWED_HOSTS", "localhost,127.0.0.1").split(",") if value.strip()]
SESSION_COOKIE = "ledgevue_session"
CSRF_COOKIE = "ledgevue_csrf"
SESSION_DAYS = 7
IMPORT_BATCH_SIZE = 1000
CATEGORIES = {
    "Salary", "Freelance", "Groceries", "Dining", "Transport", "Utilities", "Rent",
    "Subscriptions", "Shopping", "Health", "Entertainment", "Education", "Other",
}

SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
password_hash = PasswordHash.recommended()
DUMMY_PASSWORD_HASH = password_hash.hash("ledgevue-invalid-login-placeholder")


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    name: Mapped[str] = mapped_column(String(100), nullable=False)
    email: Mapped[str] = mapped_column(String(320), unique=True, nullable=False)
    password_hash: Mapped[str] = mapped_column(String(255), nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class LoginSession(Base):
    __tablename__ = "login_sessions"

    token_hash: Mapped[str] = mapped_column(String(64), primary_key=True)
    csrf_hash: Mapped[str] = mapped_column(String(64), nullable=False)
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    expires_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), nullable=False)


class TransactionRecord(Base):
    __tablename__ = "transactions"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    ref: Mapped[str] = mapped_column(String(32), nullable=False)
    date: Mapped[date] = mapped_column(Date, nullable=False)
    merchant: Mapped[str] = mapped_column(String(200), nullable=False)
    description: Mapped[str] = mapped_column(String(500), nullable=False)
    amount: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False)
    type: Mapped[str] = mapped_column(String(16), nullable=False)
    account_id: Mapped[str] = mapped_column(String(100), nullable=False)
    account_name: Mapped[str] = mapped_column(String(100), nullable=False)
    category: Mapped[str] = mapped_column(String(40), nullable=False)
    reviewed: Mapped[bool] = mapped_column(Boolean, default=False, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class BudgetRecord(Base):
    __tablename__ = "budgets"
    __table_args__ = (UniqueConstraint("user_id", "category", name="uq_budgets_user_category"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    category: Mapped[str] = mapped_column(String(40), nullable=False)
    limit: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False)


class FinancialScenarioRecord(Base):
    __tablename__ = "financial_scenarios"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    kind: Mapped[str] = mapped_column(String(20), nullable=False)
    title: Mapped[str] = mapped_column(String(120), nullable=False)
    data: Mapped[dict[str, object]] = mapped_column(JSON, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class LifeGoalRecord(Base):
    __tablename__ = "life_goals"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    title: Mapped[str] = mapped_column(String(120), nullable=False)
    target_amount: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False)
    current_amount: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False, default=Decimal("0"))
    due_date: Mapped[date] = mapped_column(Date, nullable=False)
    priority: Mapped[str] = mapped_column(String(20), nullable=False, default="normal")
    status: Mapped[str] = mapped_column(String(20), nullable=False, default="active")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class GoalContributionRecord(Base):
    __tablename__ = "goal_contributions"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    goal_id: Mapped[str] = mapped_column(ForeignKey("life_goals.id", ondelete="CASCADE"), nullable=False, index=True)
    amount: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False)
    contributed_at: Mapped[date] = mapped_column(Date, nullable=False)
    note: Mapped[str] = mapped_column(String(250), nullable=False, default="")


class FinancialRuleRecord(Base):
    __tablename__ = "financial_rules"

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    name: Mapped[str] = mapped_column(String(120), nullable=False)
    condition_type: Mapped[str] = mapped_column(String(32), nullable=False)
    category: Mapped[str | None] = mapped_column(String(40), nullable=True)
    threshold: Mapped[Decimal] = mapped_column(Numeric(14, 2), nullable=False)
    period: Mapped[str] = mapped_column(String(20), nullable=False, default="calendar_month")
    action: Mapped[str] = mapped_column(String(20), nullable=False, default="warn")
    active: Mapped[bool] = mapped_column(Boolean, nullable=False, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class RuleTriggerRecord(Base):
    __tablename__ = "financial_rule_triggers"
    __table_args__ = (UniqueConstraint("rule_id", "fingerprint", name="uq_rule_trigger_fingerprint"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    rule_id: Mapped[str] = mapped_column(ForeignKey("financial_rules.id", ondelete="CASCADE"), nullable=False, index=True)
    fingerprint: Mapped[str] = mapped_column(String(64), nullable=False)
    explanation: Mapped[str] = mapped_column(Text, nullable=False)
    evidence: Mapped[list[object]] = mapped_column(JSON, nullable=False, default=list)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), nullable=False)


class PriceChangeReviewRecord(Base):
    __tablename__ = "price_change_reviews"
    __table_args__ = (UniqueConstraint("user_id", "observation_key", name="uq_price_review_observation"),)

    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), nullable=False, index=True)
    observation_key: Mapped[str] = mapped_column(String(64), nullable=False)
    status: Mapped[str] = mapped_column(String(20), nullable=False)
    corrected_amount: Mapped[Decimal | None] = mapped_column(Numeric(14, 2), nullable=True)
    note: Mapped[str] = mapped_column(String(250), nullable=False, default="")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc), nullable=False)


class RegisterPayload(BaseModel):
    name: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=100)]
    email: EmailStr
    password: Annotated[str, StringConstraints(min_length=12, max_length=128)]


class LoginPayload(BaseModel):
    email: EmailStr
    password: Annotated[str, StringConstraints(min_length=1, max_length=128)]


class BudgetPayload(BaseModel):
    category: Annotated[str, StringConstraints(min_length=1, max_length=40)]
    limit: Decimal = Field(gt=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)


class CategoryPayload(BaseModel):
    category: Annotated[str, StringConstraints(min_length=1, max_length=40)]


class PurchaseScenarioPayload(BaseModel):
    title: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=120)] = "Planned purchase"
    amount: Decimal = Field(gt=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    category: Annotated[str, StringConstraints(min_length=1, max_length=40)] = "Other"
    purchase_date: date
    horizon_months: int = Field(ge=1, le=24)
    minimum_buffer: Decimal = Field(ge=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    save: bool = False

    @field_validator("purchase_date")
    @classmethod
    def purchase_date_is_not_in_the_past(cls, value: date) -> date:
        if value < date.today():
            raise ValueError("purchase date cannot be in the past")
        return value


class ReplayScenarioPayload(BaseModel):
    title: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=120)]
    transaction_id: str | None = None
    category: str | None = None
    alternative_amount: Decimal = Field(ge=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    monthly_saving: Decimal = Field(default=Decimal("0"), ge=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    months: int = Field(ge=1, le=24)
    save: bool = True


class GoalPayload(BaseModel):
    title: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=120)]
    target_amount: Decimal = Field(gt=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    current_amount: Decimal = Field(default=Decimal("0"), ge=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    due_date: date
    priority: Literal["low", "normal", "high"] = "normal"
    status: Literal["active", "paused", "complete"] = "active"


class GoalContributionPayload(BaseModel):
    amount: Decimal = Field(gt=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    contributed_at: date = Field(default_factory=date.today)
    note: Annotated[str, StringConstraints(max_length=250)] = ""


class RulePayload(BaseModel):
    name: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=120)]
    condition_type: Literal["transaction_amount", "category_spend", "recurring_charge", "projected_buffer", "own_transfer", "goal_behind"]
    category: str | None = None
    threshold: Decimal = Field(ge=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    period: Literal["calendar_month", "all_history"] = "calendar_month"
    action: Literal["flag", "warn", "insight"] = "warn"
    active: bool = True


if DATABASE_URL.startswith("sqlite"):
    Base.metadata.create_all(bind=engine)


class PriceReviewPayload(BaseModel):
    status: Literal["confirmed", "dismissed", "corrected"]
    corrected_amount: Decimal | None = Field(default=None, gt=0, le=Decimal("999999999999.99"), max_digits=14, decimal_places=2)
    note: Annotated[str, StringConstraints(max_length=250)] = ""


def digest(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def as_utc_datetime(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def check_origin(request: Request) -> None:
    origin = request.headers.get("origin", "").rstrip("/")
    if not origin or origin not in APP_ORIGINS:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Request origin is not allowed")


def check_csrf(request: Request, db: Session) -> LoginSession | None:
    check_origin(request)
    cookie_token = request.cookies.get(CSRF_COOKIE, "")
    header_token = request.headers.get("x-csrf-token", "")
    if len(cookie_token) < 32 or not hmac.compare_digest(cookie_token, header_token):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="CSRF validation failed")

    raw_session = request.cookies.get(SESSION_COOKIE)
    if not raw_session:
        return None
    session = db.get(LoginSession, digest(raw_session))
    if session is None or as_utc_datetime(session.expires_at) <= datetime.now(timezone.utc):
        return None
    if not hmac.compare_digest(session.csrf_hash, digest(cookie_token)):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="CSRF validation failed")
    return session


def current_user(request: Request, db: Session) -> User:
    raw_session = request.cookies.get(SESSION_COOKIE)
    if not raw_session:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    session = db.get(LoginSession, digest(raw_session))
    if session is None or as_utc_datetime(session.expires_at) <= datetime.now(timezone.utc):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    user = db.get(User, session.user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    return user


def csrf_user(request: Request, db: Session) -> User:
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    user = db.get(User, session.user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    return user


def issue_session(response: Response, db: Session, user: User) -> str:
    session_token = secrets.token_urlsafe(32)
    csrf_token = secrets.token_urlsafe(32)
    db.add(LoginSession(
        token_hash=digest(session_token),
        csrf_hash=digest(csrf_token),
        user_id=user.id,
        expires_at=datetime.now(timezone.utc) + timedelta(days=SESSION_DAYS),
    ))
    db.commit()
    response.set_cookie(SESSION_COOKIE, session_token, max_age=SESSION_DAYS * 86400, httponly=True, secure=COOKIE_SECURE, samesite="lax", path="/api")
    response.set_cookie(CSRF_COOKIE, csrf_token, max_age=SESSION_DAYS * 86400, httponly=False, secure=COOKIE_SECURE, samesite="lax", path="/")
    return csrf_token


def user_payload(user: User) -> dict[str, str]:
    return {"id": user.id, "name": user.name, "email": user.email}


def parse_date(value: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    candidate = str(value).strip()
    try:
        return datetime.fromisoformat(candidate.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for pattern in (
        "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y",
        "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%m/%d/%y",
    ):
        try:
            return datetime.strptime(candidate, pattern).date()
        except ValueError:
            continue
    raise ValueError("date must use YYYY-MM-DD or a common day/month format")


def clean_header(value: str) -> str:
    return "".join(character for character in str(value).casefold() if character.isalnum())


def first_column(headers: dict[str, int], names: tuple[str, ...]) -> int | None:
    return next((headers[name] for name in names if name in headers), None)


def parse_transaction_rows(rows: Iterable[list[object]]) -> Iterator[dict[str, object]]:
    date_names = (
        "date", "datetime", "timestamp", "transactiondate", "transactiondatetime",
        "transactiondateandtime", "transactiontimestamp", "posteddate", "postingdate",
        "valuedate", "effectivedate", "transdate", "trandate", "txndate",
        "transactiondatetimeutc", "dateoftransaction",
    )
    merchant_names = (
        "merchant", "merchantname", "payee", "payeename", "name", "narration",
        "narrative", "description", "transactiondescription", "details", "particulars",
        "transactiondetails", "transaction", "counterparty", "counterpartyname",
        "recipient", "recipientname", "vendor", "beneficiary", "descriptionmerchant", "remarks",
        "transactionremarks", "transactionparticulars", "transactionnarration",
        "merchantdescription", "merchantid", "merchantnumber",
    )
    amount_names = (
        "amount", "transactionamount", "amountinr", "netamount", "netamountinr",
        "txnamount", "amt", "value", "transactionvalue", "amountspent",
    )
    debit_names = (
        "debit", "debits", "debitamount", "debitamountinr", "debitinr",
        "withdrawal", "withdrawals", "withdrawalamount", "withdrawalamt",
        "debitamt", "dramount", "spent", "paid",
    )
    credit_names = (
        "credit", "credits", "creditamount", "creditamountinr", "creditinr",
        "deposit", "deposits", "depositamount", "depositamt", "creditamt",
        "cramount", "received",
    )

    rows = iter(rows)
    header_rows: list[list[object]] = []
    header_index = -1
    headers: dict[str, int] = {}
    for index, row in enumerate(rows):
        if index >= 25:
            break
        header_rows.append(row)
        candidate = {
            clean_header(value): column
            for column, value in enumerate(row)
            if value is not None
        }
        has_amount = any(name in candidate for name in amount_names + debit_names + credit_names)
        if (
            first_column(candidate, date_names) is not None
            and first_column(candidate, merchant_names) is not None
            and has_amount
        ):
            header_index = index
            headers = candidate
            break
    if header_index < 0:
        candidates = [
            ", ".join(str(value).strip() for value in row if value is not None and str(value).strip())
            for row in header_rows
            if any(value is not None and str(value).strip() for value in row)
        ]
        detected = max(candidates, key=len, default="none")[:180]
        raise ValueError(
            "Could not identify transaction columns. Include Date, Description, and Amount, "
            "or Date, Description, Debit, and Credit. Headers detected: " + detected
        )

    date_col = first_column(headers, date_names)
    merchant_col = first_column(headers, merchant_names)
    amount_col = first_column(headers, amount_names)
    debit_col = first_column(headers, debit_names)
    credit_col = first_column(headers, credit_names)
    description_col = first_column(headers, ("description", "memo", "narration", "details", "particulars"))
    type_col = first_column(headers, ("type", "transactiontype", "drcr", "debitcredit"))
    category_col = first_column(headers, ("category", "transactioncategory"))
    account_col = first_column(headers, ("account", "accountname", "accountid"))
    merchant_id_col = first_column(headers, ("merchantid", "merchantnumber"))
    city_col = first_column(headers, ("city", "merchantcity", "locationcity"))
    state_col = first_column(headers, ("state", "merchantstate", "region"))
    zip_col = first_column(headers, ("zip", "zipcode", "postalcode", "postcode"))
    mcc_col = first_column(headers, ("mcc", "merchantcategorycode"))
    channel_col = first_column(headers, ("usechip", "channel", "entrymode", "paymentmethod"))
    card_id_col = first_column(headers, ("cardid", "accountnumber"))
    if account_col is None:
        account_col = card_id_col
    source_id_col = first_column(headers, ("id", "transactionid", "transactionreference"))

    def cell(row: list[object], column: int | None) -> object:
        return row[column] if column is not None and column < len(row) else ""

    def text(value: object) -> str:
        return "" if value is None else str(value).strip()

    def amount_value(value: object) -> Decimal:
        raw = text(value)
        negative = raw.startswith("(") and raw.endswith(")")
        raw = raw.strip("() ").replace(",", "")
        raw = raw.replace("₹", "").replace("$", "")
        raw = raw.casefold().replace("inr", "").replace("rs.", "").replace("rs", "")
        suffix = raw.casefold()[-2:]
        if suffix in {"dr", "cr"}:
            raw = raw[:-2].strip()
            negative = negative or suffix == "dr"
        value_decimal = Decimal(raw)
        return -abs(value_decimal) if negative else value_decimal

    parsed_count = 0
    summary_rows = {
        "total", "grandtotal", "subtotal", "openingbalance", "closingbalance",
        "balancebroughtforward", "balancecarriedforward", "broughtforward",
        "carriedforward", "endofstatement", "endofreport",
    }
    data_rows = chain(header_rows[header_index + 1 :], rows)
    for row_index, row in enumerate(data_rows, start=header_index + 2):
        if not any(text(value) for value in row):
            continue
        repeated_date = clean_header(cell(row, date_col))
        repeated_merchant = clean_header(cell(row, merchant_col))
        if repeated_date in date_names and repeated_merchant in merchant_names:
            continue
        if repeated_date in summary_rows or repeated_merchant in summary_rows:
            continue
        date_value = text(cell(row, date_col))
        merchant = text(cell(row, merchant_col))
        if not date_value and not merchant:
            continue
        try:
            if not merchant or len(merchant) > 200:
                raise ValueError("merchant is required and must be under 200 characters")

            location_parts = [text(cell(row, city_col)), text(cell(row, state_col))]
            postal_code = text(cell(row, zip_col)).removesuffix(".0")
            location = ", ".join(part for part in location_parts if part)
            if postal_code:
                location = f"{location} {postal_code}".strip()
            merchant_id = text(cell(row, merchant_id_col))
            if merchant_id_col is not None and merchant_col == merchant_id_col:
                merchant = f"Merchant {merchant_id}"
                if location:
                    merchant = f"{merchant} ({location})"
            if not merchant or len(merchant) > 200:
                raise ValueError("merchant details are missing or too long")

            raw_amount = text(cell(row, amount_col)) if amount_col is not None else ""
            raw_debit = text(cell(row, debit_col)) if debit_col is not None else ""
            raw_credit = text(cell(row, credit_col)) if credit_col is not None else ""
            if raw_amount:
                amount = amount_value(raw_amount)
            elif raw_debit and amount_value(raw_debit) != 0:
                amount = -abs(amount_value(raw_debit))
            elif raw_credit and amount_value(raw_credit) != 0:
                amount = abs(amount_value(raw_credit))
            else:
                raise ValueError("amount, debit, or credit is required")
            if not amount.is_finite() or abs(amount) > Decimal("999999999999.99"):
                raise ValueError("amount must be non-zero and within the supported range")
            if amount == 0:
                continue

            raw_type = text(cell(row, type_col)).casefold() if type_col is not None else ""
            if raw_type in {"income", "credit", "deposit", "cr"}:
                tx_type = "income"
            elif raw_type in {"refund", "reversal"}:
                tx_type = "refund"
            elif raw_type in {"transfer", "internal transfer"}:
                tx_type = "transfer"
            elif raw_type in {"expense", "debit", "purchase", "dr"}:
                tx_type = "expense"
            elif raw_type == "":
                if raw_debit and amount_value(raw_debit) != 0 and not raw_amount:
                    tx_type = "expense"
                elif raw_credit and amount_value(raw_credit) != 0 and not raw_amount:
                    tx_type = "income"
                else:
                    tx_type = "income" if amount > 0 else "expense"
            else:
                raise ValueError("type must be income, expense, refund, or transfer")

            category = text(cell(row, category_col)).title() if category_col is not None else "Other"
            if category not in CATEGORIES:
                category = "Other"
            account_name = (text(cell(row, account_col)) or "Imported account")[:100]
            if card_id_col is not None and account_col == card_id_col:
                account_name = f"Card {account_name}"[:100]
            account_id = str(uuid.uuid5(uuid.NAMESPACE_URL, account_name.casefold()))
            description_value = text(cell(row, description_col))
            if not description_value:
                description_parts = [text(cell(row, channel_col))]
                mcc = text(cell(row, mcc_col))
                if mcc:
                    description_parts.append(f"MCC {mcc}")
                if location:
                    description_parts.append(location)
                description_value = " · ".join(part for part in description_parts if part)
            description = (description_value or merchant)[:500]
            source_id = text(cell(row, source_id_col))
            parsed_count += 1
            yield {
                "id": str(uuid.uuid4()),
                "ref": (f"IMP-{source_id}" if source_id else f"IMP-{uuid.uuid4().hex[:12].upper()}")[:32],
                "date": parse_date(cell(row, date_col)),
                "merchant": merchant,
                "description": description,
                "amount": abs(amount),
                "type": tx_type,
                "account_id": account_id,
                "account_name": account_name,
                "category": category,
            }
        except (InvalidOperation, ValueError) as error:
            raise ValueError(f"Row {row_index}: {error}") from error
    if parsed_count == 0:
        raise ValueError("File contains no transaction rows")


def iter_csv_rows(file_obj: BinaryIO) -> Iterator[dict[str, object]]:
    try:
        file_obj.seek(0)
        signature = file_obj.read(2)
        file_obj.seek(0)
        encoding = "utf-16" if signature in (b"\xff\xfe", b"\xfe\xff") else "utf-8-sig"
        text_stream = io.TextIOWrapper(file_obj, encoding=encoding, newline="")
        sample_lines = []
        try:
            for _ in range(25):
                line = text_stream.readline()
                if not line:
                    break
                sample_lines.append(line)
            if not sample_lines:
                raise ValueError("CSV contains no rows")
            declared_delimiter = None
            if sample_lines[0].casefold().startswith("sep="):
                declared_delimiter = sample_lines.pop(0)[4:5]
            sample = "".join(sample_lines)
            delimiters = ",;\t|"
            if declared_delimiter and declared_delimiter in delimiters:
                delimiter = declared_delimiter
            else:
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=delimiters).delimiter
                except csv.Error:
                    lines = [line for line in sample.splitlines() if line.strip()]
                    if not lines:
                        raise ValueError("CSV contains no rows")
                    header_candidate = max(
                        lines, key=lambda line: max(line.count(value) for value in delimiters)
                    )
                    delimiter = max(delimiters, key=header_candidate.count)
            rows = csv.reader(chain(sample_lines, text_stream), delimiter=delimiter)
            yield from parse_transaction_rows(rows)
        finally:
            text_stream.detach()
    except UnicodeDecodeError as error:
        raise ValueError("CSV must be UTF-8 or UTF-16 encoded") from error
    except csv.Error as error:
        raise ValueError("CSV format could not be read") from error


def parse_csv_rows(payload: bytes) -> list[dict[str, object]]:
    return list(iter_csv_rows(io.BytesIO(payload)))


def iter_uploaded_rows(file_obj: BinaryIO, filename: str) -> Iterator[dict[str, object]]:
    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if extension == "csv":
        yield from iter_csv_rows(file_obj)
        return
    if extension == "xlsx":
        try:
            workbook = load_workbook(file_obj, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            raise ValueError("Excel file could not be read; save it as .xlsx and try again") from error
        try:
            if not workbook.worksheets:
                raise ValueError("Excel workbook has no worksheets")
            rows = (list(row) for row in workbook.worksheets[0].iter_rows(values_only=True))
            yield from parse_transaction_rows(rows)
        finally:
            workbook.close()
        return
    if extension == "xls":
        workbook = None
        try:
            workbook = xlrd.open_workbook(file_contents=file_obj.read(), on_demand=True)
            sheet = workbook.sheet_by_index(0)
            rows: list[list[object]] = []
            for row_index in range(sheet.nrows):
                values = []
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    values.append(value)
                rows.append(values)
            yield from parse_transaction_rows(rows)
        except (IndexError, xlrd.biffh.XLRDError, ValueError) as error:
            raise ValueError("Excel file could not be read; save it as .xls and try again") from error
        finally:
            if workbook is not None:
                workbook.release_resources()
        return
    raise ValueError("Supported files are CSV, XLSX, and XLS")


def parse_uploaded_rows(payload: bytes, filename: str) -> list[dict[str, object]]:
    return list(iter_uploaded_rows(io.BytesIO(payload), filename))


def import_file_rows(user_id: str, file_obj: BinaryIO, filename: str) -> int:
    db = SessionLocal()
    imported = 0
    try:
        rows = iter_uploaded_rows(file_obj, filename)
        while batch := list(islice(rows, IMPORT_BATCH_SIZE)):
            db.bulk_insert_mappings(
                TransactionRecord,
                [{"user_id": user_id, **row} for row in batch],
            )
            imported += len(batch)
        db.commit()
        return imported
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def _money(value: Decimal | int | float | str) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.01")))


def _detect_recurring_for_user(rows: list[TransactionRecord]) -> list[dict[str, object]]:
    by_key: dict[str, list[TransactionRecord]] = {}
    for row in rows:
        if row.type != "expense":
            continue
        key = f"{row.merchant.strip().lower()}|{row.category}"
        by_key.setdefault(key, []).append(row)

    recurring: list[dict[str, object]] = []
    for entries in by_key.values():
        if len(entries) < 3:
            continue
        ordered = sorted(entries, key=lambda item: item.date)
        gaps = [
            (ordered[index + 1].date - ordered[index].date).days
            for index in range(len(ordered) - 1)
        ]
        if not gaps:
            continue
        avg_gap = sum(gaps) / len(gaps)
        frequency = "monthly" if 24 <= avg_gap <= 40 else "quarterly" if 78 <= avg_gap <= 102 else None
        if frequency is None:
            continue
        if any(abs(days - avg_gap) > (8 if frequency == "monthly" else 14) for days in gaps):
            continue
        amounts = [float(item.amount) for item in ordered]
        mean_amount = sum(amounts) / len(amounts)
        if any(abs(amount - mean_amount) > max(mean_amount * 0.15, 1) for amount in amounts):
            continue
        next_date = ordered[-1].date + timedelta(days=30 if frequency == "monthly" else 90)
        recurring.append({
            "id": ordered[0].id,
            "merchant": ordered[0].merchant,
            "category": ordered[0].category,
            "amount": round(mean_amount),
            "frequency": frequency,
            "lastPaid": ordered[-1].date.isoformat(),
            "nextExpected": next_date.isoformat(),
            "confidence": "detected",
        })
    return sorted(recurring, key=lambda item: item["amount"], reverse=True)


def _merchant_key(value: str) -> str:
    ignored = {"subscription", "payment", "online", "txn", "transaction", "limited", "ltd", "inc"}
    return " ".join(token for token in re.findall(r"[a-z0-9]+", value.casefold()) if token not in ignored)


def _detect_price_creep(rows: list[TransactionRecord]) -> list[dict[str, object]]:
    by_merchant: dict[str, list[TransactionRecord]] = {}
    for row in rows:
        if row.type == "expense" and row.amount > 0:
            key = _merchant_key(row.merchant)
            if key:
                by_merchant.setdefault(key, []).append(row)

    alerts: list[dict[str, object]] = []

    def frequency_for_gap(gap: float) -> str | None:
        if 5 <= gap <= 9:
            return "weekly"
        if 24 <= gap <= 40:
            return "monthly"
        if 78 <= gap <= 102:
            return "quarterly"
        if 340 <= gap <= 390:
            return "annual"
        return None

    for entries in by_merchant.values():
        ordered = sorted(entries, key=lambda item: item.date)
        if len(ordered) < 3:
            continue
        gaps = [(ordered[index + 1].date - ordered[index].date).days for index in range(len(ordered) - 1)]
        average_gap = sum(gaps) / len(gaps)
        previous_frequency = frequency_for_gap(gaps[-2]) if len(gaps) > 1 else None
        latest_frequency = frequency_for_gap(gaps[-1])
        if previous_frequency and latest_frequency and previous_frequency != latest_frequency:
            change_type = "billing_frequency_change"
        elif latest_frequency and all(frequency_for_gap(gap) == latest_frequency for gap in gaps):
            change_type = "possible_price_change"
        else:
            continue
        if change_type == "possible_price_change" and any(abs(gap - average_gap) > 14 for gap in gaps):
            continue
        previous = float(ordered[-2].amount)
        current = float(ordered[-1].amount)
        if current <= 0 or previous <= 0:
            continue
        change = current - previous
        percentage = (change / previous) * 100
        if change_type == "possible_price_change" and abs(change) < 50 and abs(percentage) < 10:
            continue
        alerts.append({
            "merchant": ordered[-1].merchant,
            "previousAmount": round(previous),
            "newAmount": round(current),
            "change": round(change),
            "percentage": round(percentage, 1),
            "period": f"{ordered[-2].date.isoformat()} → {ordered[-1].date.isoformat()}",
            "evidenceDates": [ordered[-2].date.isoformat(), ordered[-1].date.isoformat()],
            "evidenceTransactionIds": [ordered[-2].id, ordered[-1].id],
            "frequency": latest_frequency,
            "previousFrequency": previous_frequency,
            "changeType": change_type,
            "confidence": "possible",
            "observationKey": digest(f"{ordered[-1].user_id}|{_merchant_key(ordered[-1].merchant)}|{ordered[-2].date}|{ordered[-1].date}|{previous}|{current}"),
        })
    return sorted(alerts, key=lambda item: abs(float(item["change"])), reverse=True)


def _simulate_purchase_impact(
    rows: list[TransactionRecord],
    planned_purchase: Decimal,
    purchase_date: date,
    horizon_months: int,
    minimum_buffer: Decimal,
    goals: list[LifeGoalRecord] | None = None,
) -> dict[str, object]:
    months = sorted({(row.date.year, row.date.month) for row in rows})[-3:]
    if not rows or not any(row.type == "income" for row in rows) or len(months) < 2:
        return {
            "plannedPurchase": _money(planned_purchase),
            "purchaseDate": purchase_date.isoformat(),
            "horizonMonths": horizon_months,
            "projectedBuffer": None,
            "minimumBuffer": _money(minimum_buffer),
            "risk": "Insufficient Data",
            "signal": "Insufficient Data",
            "upcomingBills": 0,
            "availableFunds": None,
            "explanation": ["At least two months of records including income are needed for a cash-flow estimate.", "Account balances, planned savings, and bill due dates are not available in imported transactions, so this is not a verified cash balance."],
            "missingInputs": ["Current account balances", "User-entered bill due dates"],
        }

    monthly_net: dict[tuple[int, int], Decimal] = {key: Decimal("0") for key in months}
    for item in rows:
        key = (item.date.year, item.date.month)
        if key not in monthly_net:
            continue
        if item.type == "income":
            monthly_net[key] += item.amount
        elif item.type == "expense":
            monthly_net[key] -= item.amount
        elif item.type == "refund":
            monthly_net[key] += item.amount
    average_monthly_net = sum(monthly_net.values(), Decimal("0")) / len(monthly_net)
    recurring = _detect_recurring_for_user(rows)
    expected_bills = Decimal("0")
    recurring_monthly = Decimal("0")
    horizon_end = purchase_date + timedelta(days=horizon_months * 31)
    for item in recurring:
        amount = Decimal(str(item["amount"]))
        interval = 90 if item["frequency"] == "quarterly" else 30
        recurring_monthly += amount / (Decimal("3") if interval == 90 else Decimal("1"))
        next_expected = date.fromisoformat(str(item["nextExpected"]))
        while next_expected < purchase_date:
            next_expected += timedelta(days=interval)
        while next_expected <= horizon_end:
            expected_bills += amount
            next_expected += timedelta(days=interval)
    planned_savings = Decimal("0")
    for goal in goals or []:
        if goal.status != "active":
            continue
        months_left = max(1, (goal.due_date.year - date.today().year) * 12 + goal.due_date.month - date.today().month)
        planned_savings += max(Decimal("0"), goal.target_amount - goal.current_amount) / months_left
    available_funds = (average_monthly_net + recurring_monthly) * horizon_months - expected_bills - planned_savings * horizon_months
    projected_buffer = available_funds - planned_purchase
    signal = "Buffer at Risk" if projected_buffer < minimum_buffer else "Within Plan"

    return {
        "plannedPurchase": _money(planned_purchase),
        "purchaseDate": purchase_date.isoformat(),
        "horizonMonths": horizon_months,
        "projectedBuffer": _money(projected_buffer),
        "minimumBuffer": _money(minimum_buffer),
        "risk": "Insufficient Data",
        "signal": signal,
        "upcomingBills": _money(expected_bills),
        "availableFunds": _money(available_funds),
        "explanation": [
            f"Estimated net cash flow across {horizon_months} month(s), based on {len(monthly_net)} imported calendar months and adjusted for recurring charges: ₹{_money(available_funds):,.2f}.",
            f"Detected recurring payments due in the selected period: about ₹{_money(expected_bills):,.2f}.",
            f"Active goal savings pace across this period: about ₹{_money(planned_savings * horizon_months):,.2f}.",
            f"After this purchase, the estimated net-flow remainder is ₹{_money(projected_buffer):,.2f}, compared with your ₹{_money(minimum_buffer):,.2f} preference. Indicative signal: {signal}.",
            "This excludes starting account balances and unrecorded bills or savings. Treat it as a scenario comparison, not an account-balance forecast.",
        ],
        "missingInputs": ["Current account balances", "User-entered bill due dates"],
    }


def _goal_payload(goal: LifeGoalRecord) -> dict[str, object]:
    remaining = max(Decimal("0"), goal.target_amount - goal.current_amount)
    today = date.today()
    days_remaining = max(0, (goal.due_date - today).days)
    months_remaining = ceil(days_remaining / 30.4375)
    month_count = max(1, months_remaining)
    return {
        "id": goal.id,
        "title": goal.title,
        "targetAmount": _money(goal.target_amount),
        "currentAmount": _money(goal.current_amount),
        "remainingAmount": _money(remaining),
        "dueDate": goal.due_date.isoformat(),
        "monthlyTarget": _money(remaining / month_count),
        "monthsRemaining": months_remaining,
        "progress": min(100, round(float(goal.current_amount / goal.target_amount * 100))) if goal.target_amount else 0,
        "priority": goal.priority,
        "status": goal.status,
    }


def _rule_payload(rule: FinancialRuleRecord) -> dict[str, object]:
    return {
        "id": rule.id,
        "name": rule.name,
        "conditionType": rule.condition_type,
        "category": rule.category,
        "threshold": _money(rule.threshold),
        "period": rule.period,
        "action": rule.action,
        "active": rule.active,
    }


def _scenario_payload(scenario: FinancialScenarioRecord) -> dict[str, object]:
    return {
        "id": scenario.id,
        "kind": scenario.kind,
        "title": scenario.title,
        "data": scenario.data,
        "createdAt": scenario.created_at.isoformat(),
    }


def _trigger_payload(trigger: RuleTriggerRecord) -> dict[str, object]:
    return {
        "id": trigger.id,
        "ruleId": trigger.rule_id,
        "explanation": trigger.explanation,
        "evidence": trigger.evidence,
        "createdAt": trigger.created_at.isoformat(),
    }


def _evaluate_rules(db: Session, user_id: str, rows: list[TransactionRecord]) -> None:
    active_rules = db.scalars(select(FinancialRuleRecord).where(FinancialRuleRecord.user_id == user_id, FinancialRuleRecord.active.is_(True))).all()
    today = date.today()
    month_start = today.replace(day=1)
    for rule in active_rules:
        period_start = month_start if rule.period == "calendar_month" else date.min
        matching: list[TransactionRecord] = []
        value = Decimal("0")
        explanation = ""
        evidence: list[object] = []
        if rule.condition_type == "transaction_amount":
            matching = [item for item in rows if item.amount >= rule.threshold and item.date >= period_start]
            if matching:
                value = max(item.amount for item in matching)
                explanation = f"A transaction of {_money(value):,.2f} met or exceeded this rule's {_money(rule.threshold):,.2f} threshold."
                evidence = [{"transactionId": item.id, "merchant": item.merchant, "amount": _money(item.amount), "date": item.date.isoformat()} for item in matching[-10:]]
        elif rule.condition_type == "category_spend":
            matching = [item for item in rows if item.type == "expense" and item.category == rule.category and item.date >= period_start]
            value = sum((item.amount for item in matching), Decimal("0"))
            if value > rule.threshold:
                explanation = f"{rule.category} expenses total {_money(value):,.2f}, above the {_money(rule.threshold):,.2f} threshold."
                evidence = [{"transactionId": item.id, "amount": _money(item.amount), "date": item.date.isoformat()} for item in matching[-20:]]
        elif rule.condition_type == "recurring_charge":
            recurring = _detect_recurring_for_user(rows)
            matched_recurring = [item for item in recurring if Decimal(str(item["amount"])) >= rule.threshold and (rule.category is None or item["category"] == rule.category)]
            if matched_recurring:
                value = max(Decimal(str(item["amount"])) for item in matched_recurring)
                explanation = f"A detected recurring payment meets or exceeds {_money(rule.threshold):,.2f}. Review the pattern before treating it as a confirmed subscription."
                evidence = matched_recurring
        elif rule.condition_type == "projected_buffer":
            months = sorted({(item.date.year, item.date.month) for item in rows})[-3:]
            if len(months) >= 2:
                monthly_net = {key: Decimal("0") for key in months}
                for item in rows:
                    key = (item.date.year, item.date.month)
                    if key not in monthly_net:
                        continue
                    if item.type == "income":
                        monthly_net[key] += item.amount
                    elif item.type == "expense":
                        monthly_net[key] -= item.amount
                    elif item.type == "refund":
                        monthly_net[key] += item.amount
                value = sum(monthly_net.values(), Decimal("0")) / len(months)
                if value < rule.threshold:
                    explanation = f"Recent average monthly net cash flow is {_money(value):,.2f}, below the configured {_money(rule.threshold):,.2f} threshold. This is not an account balance."
                    evidence = [{"month": f"{year:04d}-{month:02d}", "netCashFlow": _money(net)} for (year, month), net in monthly_net.items()]
        elif rule.condition_type == "own_transfer":
            matching = [item for item in rows if item.type == "transfer" and item.amount >= rule.threshold and item.date >= period_start]
            if matching:
                value = max(item.amount for item in matching)
                explanation = f"{len(matching)} transaction(s) classified as transfers met the configured threshold. Review the source records; the import does not independently verify account ownership."
                evidence = [{"transactionId": item.id, "merchant": item.merchant, "amount": _money(item.amount), "date": item.date.isoformat()} for item in matching[-10:]]
        elif rule.condition_type == "goal_behind":
            goals = db.scalars(select(LifeGoalRecord).where(LifeGoalRecord.user_id == user_id, LifeGoalRecord.status == "active")).all()
            behind = []
            for goal in goals:
                total_days = max(1, (goal.due_date - goal.created_at.date()).days)
                elapsed_days = max(0, (today - goal.created_at.date()).days)
                expected_saved = goal.target_amount * min(1, elapsed_days / total_days)
                if goal.current_amount + rule.threshold < expected_saved:
                    behind.append({"goalId": goal.id, "title": goal.title, "currentAmount": _money(goal.current_amount), "expectedByNow": _money(expected_saved), "targetDate": goal.due_date.isoformat()})
            if behind:
                explanation = f"{len(behind)} active goal(s) are behind their straight-line savings pace by more than the {_money(rule.threshold):,.2f} tolerance."
                evidence = behind
        if not explanation:
            continue
        period = today.strftime("%Y-%m") if rule.period == "calendar_month" else "all-history"
        if rule.condition_type in {"category_spend", "projected_buffer", "goal_behind"}:
            fingerprint = digest(f"{rule.id}|{period}|triggered")
        else:
            fingerprint = digest(f"{rule.id}|{period}|{json.dumps(evidence, sort_keys=True, default=str)}")
        existing = db.scalar(select(RuleTriggerRecord.id).where(RuleTriggerRecord.rule_id == rule.id, RuleTriggerRecord.fingerprint == fingerprint))
        if existing is None:
            db.add(RuleTriggerRecord(user_id=user_id, rule_id=rule.id, fingerprint=fingerprint, explanation=explanation, evidence=evidence))
    db.commit()


def _build_financial_intelligence(rows: list[TransactionRecord], budgets: list[BudgetRecord], db: Session, user_id: str) -> dict[str, object]:
    recurring = _detect_recurring_for_user(rows)
    price_creep = _detect_price_creep(rows)
    reviews = {row.observation_key: row for row in db.scalars(select(PriceChangeReviewRecord).where(PriceChangeReviewRecord.user_id == user_id)).all()}
    for item in price_creep:
        review = reviews.get(str(item["observationKey"]))
        item["reviewStatus"] = review.status if review else "pending"
        item["correctedAmount"] = _money(review.corrected_amount) if review and review.corrected_amount is not None else None
        if review and review.status == "corrected" and review.corrected_amount is not None:
            item["newAmount"] = _money(review.corrected_amount)
            item["change"] = _money(review.corrected_amount - Decimal(str(item["previousAmount"])))
            item["percentage"] = round(float((review.corrected_amount - Decimal(str(item["previousAmount"]))) / Decimal(str(item["previousAmount"])) * 100), 1)
    goals = [_goal_payload(goal) for goal in db.scalars(select(LifeGoalRecord).where(LifeGoalRecord.user_id == user_id).order_by(LifeGoalRecord.due_date)).all()]
    rules = [_rule_payload(rule) for rule in db.scalars(select(FinancialRuleRecord).where(FinancialRuleRecord.user_id == user_id).order_by(FinancialRuleRecord.created_at.desc())).all()]
    recent_months = sorted({(row.date.year, row.date.month) for row in rows})[-3:]
    recent_rows = [row for row in rows if (row.date.year, row.date.month) in recent_months]
    monthly_income = sum((row.amount for row in recent_rows if row.type == "income"), Decimal("0")) / max(1, len(recent_months))
    monthly_expense = sum((row.amount for row in recent_rows if row.type == "expense"), Decimal("0")) / max(1, len(recent_months))
    buffer_estimate = monthly_income - monthly_expense
    missing_inputs = not rows or len({(row.date.year, row.date.month) for row in rows}) < 2 or not any(row.type == "income" for row in rows)
    return {
        "runway": {
            "score": None if missing_inputs else max(0, min(100, round(100 * (monthly_income - monthly_expense) / max(monthly_income, Decimal("1"))))),
            "status": "Insufficient Data" if missing_inputs else "Cash-flow estimate",
            "runwayMonths": None,
            "monthlyIncome": _money(monthly_income),
            "monthlyExpense": _money(monthly_expense),
            "buffer": _money(buffer_estimate),
            "insights": ["Imported transaction cash flow is not the same as an account balance.", "Add at least two months of income and expense records for a net-flow comparison." if missing_inputs else "Estimate uses the most recent three calendar months; it excludes starting account balances and unrecorded commitments."],
        },
        "priceCreep": price_creep,
        "rules": rules,
        "goals": goals,
        "recurring": recurring,
        "scenarios": [_scenario_payload(row) for row in db.scalars(select(FinancialScenarioRecord).where(FinancialScenarioRecord.user_id == user_id).order_by(FinancialScenarioRecord.created_at.desc()).limit(100)).all()],
        "ruleTriggers": [_trigger_payload(row) for row in db.scalars(select(RuleTriggerRecord).where(RuleTriggerRecord.user_id == user_id).order_by(RuleTriggerRecord.created_at.desc()).limit(100)).all()],
    }


limiter = Limiter(key_func=get_remote_address, default_limits=["120/minute"])
app = FastAPI(title="LedgeVue API", docs_url=None, redoc_url=None, openapi_url=None)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=ALLOWED_HOSTS)


@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = "default-src 'none'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/auth/csrf")
def get_csrf(response: Response, request: Request, db: Session = Depends(get_db)):
    raw_session = request.cookies.get(SESSION_COOKIE)
    session = db.get(LoginSession, digest(raw_session)) if raw_session else None
    if session and as_utc_datetime(session.expires_at) > datetime.now(timezone.utc):
        token = secrets.token_urlsafe(32)
        session.csrf_hash = digest(token)
        db.commit()
    else:
        token = secrets.token_urlsafe(32)
    response.set_cookie(CSRF_COOKIE, token, max_age=SESSION_DAYS * 86400, httponly=False, secure=COOKIE_SECURE, samesite="lax", path="/")
    return {"csrfToken": token}


@app.post("/api/auth/register")
@limiter.limit("5/minute")
def register(request: Request, payload: RegisterPayload, response: Response, db: Session = Depends(get_db)):
    check_csrf(request, db)
    user = User(name=payload.name, email=str(payload.email).lower(), password_hash=password_hash.hash(payload.password))
    db.add(user)
    try:
        db.flush()
        issue_session(response, db, user)
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="An account with this email already exists") from error
    return {"user": user_payload(user)}


@app.post("/api/auth/login")
@limiter.limit("5/minute")
def login(request: Request, payload: LoginPayload, response: Response, db: Session = Depends(get_db)):
    check_csrf(request, db)
    user = db.scalar(select(User).where(User.email == str(payload.email).lower()))
    stored_hash = user.password_hash if user else DUMMY_PASSWORD_HASH
    password_matches = password_hash.verify(payload.password, stored_hash)
    if user is None or not password_matches:
        raise HTTPException(status_code=401, detail="Email or password is incorrect")
    issue_session(response, db, user)
    return {"user": user_payload(user)}


@app.get("/api/auth/me")
def me(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    return {"user": user_payload(user)}


@app.post("/api/auth/logout")
def logout(request: Request, response: Response, db: Session = Depends(get_db)):
    check_csrf(request, db)
    raw_session = request.cookies.get(SESSION_COOKIE)
    if raw_session:
        db.query(LoginSession).filter(LoginSession.token_hash == digest(raw_session)).delete()
        db.commit()
    response.delete_cookie(SESSION_COOKIE, path="/api", httponly=True, secure=COOKIE_SECURE, samesite="lax")
    response.delete_cookie(CSRF_COOKIE, path="/", secure=COOKIE_SECURE, samesite="lax")
    return {"ok": True}


@app.get("/api/finance")
def finance(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    rows = db.scalars(
        select(TransactionRecord)
        .where(TransactionRecord.user_id == user.id)
        .order_by(TransactionRecord.date.desc(), TransactionRecord.created_at.desc())
        .limit(50000)
    ).all()
    accounts = {}
    transactions = []
    for row in rows:
        accounts[row.account_id] = {"id": row.account_id, "name": row.account_name, "kind": "bank"}
        transactions.append({
            "id": row.id, "ref": row.ref, "date": row.date.isoformat(), "merchant": row.merchant,
            "description": row.description, "amount": float(row.amount), "type": row.type,
            "accountId": row.account_id, "category": row.category, "reviewed": row.reviewed,
        })
    latest = max((row.created_at for row in rows), default=None)
    budget_rows = db.scalars(select(BudgetRecord).where(BudgetRecord.user_id == user.id)).all()
    budgets = [{"category": row.category, "limit": float(row.limit)} for row in budget_rows]
    recurring = _detect_recurring_for_user(rows)
    _evaluate_rules(db, user.id, rows)
    intelligence = _build_financial_intelligence(rows, budget_rows, db, user.id)
    return {
        "transactions": transactions,
        "accounts": list(accounts.values()),
        "budgets": budgets,
        "recurring": recurring,
        "intelligence": intelligence,
        "lastImportedAt": latest.isoformat() if latest else None,
    }


@app.get("/api/finance/intelligence")
def finance_intelligence(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    rows = db.scalars(
        select(TransactionRecord)
        .where(TransactionRecord.user_id == user.id)
        .order_by(TransactionRecord.date.desc(), TransactionRecord.created_at.desc())
        .limit(50000)
    ).all()
    budget_rows = db.scalars(select(BudgetRecord).where(BudgetRecord.user_id == user.id)).all()
    _evaluate_rules(db, user.id, rows)
    return _build_financial_intelligence(rows, budget_rows, db, user.id)


@app.post("/api/intelligence/purchase-impact")
def purchase_impact(request: Request, payload: PurchaseScenarioPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    rows = db.scalars(select(TransactionRecord).where(TransactionRecord.user_id == user.id).order_by(TransactionRecord.date)).all()
    goals = db.scalars(select(LifeGoalRecord).where(LifeGoalRecord.user_id == user.id, LifeGoalRecord.status == "active")).all()
    result = _simulate_purchase_impact(rows, payload.amount, payload.purchase_date, payload.horizon_months, payload.minimum_buffer, goals)
    result["category"] = payload.category
    if payload.save:
        scenario = FinancialScenarioRecord(user_id=user.id, kind="purchase", title=payload.title, data=result)
        db.add(scenario)
        db.commit()
        result["scenarioId"] = scenario.id
    return result


@app.get("/api/intelligence/scenarios")
def list_scenarios(request: Request, kind: Literal["purchase", "replay"] | None = None, limit: int = 50, offset: int = 0, db: Session = Depends(get_db)):
    user = current_user(request, db)
    query = select(FinancialScenarioRecord).where(FinancialScenarioRecord.user_id == user.id)
    if kind:
        query = query.where(FinancialScenarioRecord.kind == kind)
    page_size = min(100, max(1, limit))
    scenarios = db.scalars(query.order_by(FinancialScenarioRecord.created_at.desc()).offset(max(0, offset)).limit(page_size)).all()
    return {"items": [_scenario_payload(item) for item in scenarios], "limit": page_size, "offset": max(0, offset)}


@app.delete("/api/intelligence/scenarios/{scenario_id}")
def delete_scenario(scenario_id: str, request: Request, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    scenario = db.scalar(select(FinancialScenarioRecord).where(FinancialScenarioRecord.id == scenario_id, FinancialScenarioRecord.user_id == user.id))
    if scenario is None:
        raise HTTPException(status_code=404, detail="Scenario not found")
    db.delete(scenario)
    db.commit()
    return {"deleted": True}


@app.post("/api/intelligence/replays", status_code=status.HTTP_201_CREATED)
def create_replay(request: Request, payload: ReplayScenarioPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    if not payload.transaction_id and not payload.category:
        raise HTTPException(status_code=422, detail="Select a past expense or category to replay")
    transaction = None
    if payload.transaction_id:
        transaction = db.scalar(select(TransactionRecord).where(TransactionRecord.id == payload.transaction_id, TransactionRecord.user_id == user.id))
        if transaction is None or transaction.type != "expense":
            raise HTTPException(status_code=404, detail="Eligible expense transaction not found")
    all_rows = db.scalars(select(TransactionRecord).where(TransactionRecord.user_id == user.id)).all()
    if transaction:
        recorded_per_month = transaction.amount
        period_type = "one-time purchase"
        actual_total = transaction.amount
        hypothetical_total = payload.alternative_amount
    else:
        eligible = [item for item in all_rows if item.type == "expense" and item.category == payload.category]
        if not eligible:
            raise HTTPException(status_code=422, detail="No imported expenses are available for that category")
        period_keys = sorted({(item.date.year, item.date.month) for item in all_rows})[-payload.months:]
        monthly_amounts = {key: sum((item.amount for item in eligible if (item.date.year, item.date.month) == key), Decimal("0")) for key in period_keys}
        recorded_per_month = sum(monthly_amounts.values(), Decimal("0")) / max(1, len(monthly_amounts))
        period_type = "monthly category average"
        actual_total = recorded_per_month * payload.months
        hypothetical_total = payload.alternative_amount * payload.months
    contributions = payload.monthly_saving * payload.months
    potential_savings = max(Decimal("0"), actual_total - hypothetical_total) + contributions
    data = {
        "basis": period_type,
        "transactionId": transaction.id if transaction else None,
        "category": payload.category or (transaction.category if transaction else None),
        "months": payload.months,
        "historyMonthsUsed": len(period_keys) if not transaction else 1,
        "actualAmount": _money(actual_total),
        "alternativeAmount": _money(hypothetical_total),
        "monthlySaving": _money(payload.monthly_saving),
        "contributions": _money(contributions),
        "potentialSavings": _money(potential_savings),
        "timeline": [{"month": index, "contribution": _money(payload.monthly_saving), "cumulativeContribution": _money(payload.monthly_saving * index)} for index in range(1, payload.months + 1)],
        "assumptions": ["This is hypothetical and does not change past transactions or account balances.", "No interest or investment return is included.", "Category replay uses the average spend in the available selected history; missing months count as zero only through the selected comparison horizon."],
    }
    scenario = FinancialScenarioRecord(user_id=user.id, kind="replay", title=payload.title, data=data)
    db.add(scenario)
    db.commit()
    data["scenarioId"] = scenario.id
    return {"id": scenario.id, "kind": "replay", "title": scenario.title, "data": data, "createdAt": scenario.created_at.isoformat()}


@app.patch("/api/intelligence/price-changes/{observation_key}/review")
def review_price_change(observation_key: str, request: Request, payload: PriceReviewPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    if payload.status == "corrected" and payload.corrected_amount is None:
        raise HTTPException(status_code=422, detail="Enter the corrected recurring amount")
    review = db.scalar(select(PriceChangeReviewRecord).where(PriceChangeReviewRecord.user_id == user.id, PriceChangeReviewRecord.observation_key == observation_key))
    if review is None:
        review = PriceChangeReviewRecord(user_id=user.id, observation_key=observation_key, status=payload.status)
        db.add(review)
    review.status = payload.status
    review.corrected_amount = payload.corrected_amount if payload.status == "corrected" else None
    review.note = payload.note
    db.commit()
    return {"observationKey": review.observation_key, "status": review.status, "correctedAmount": _money(review.corrected_amount) if review.corrected_amount is not None else None}


@app.get("/api/intelligence/goals")
def list_goals(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    goals = db.scalars(select(LifeGoalRecord).where(LifeGoalRecord.user_id == user.id).order_by(LifeGoalRecord.due_date)).all()
    return {"items": [_goal_payload(goal) for goal in goals]}


@app.post("/api/intelligence/goals", status_code=status.HTTP_201_CREATED)
def create_goal(request: Request, payload: GoalPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    goal = LifeGoalRecord(user_id=user.id, title=payload.title, target_amount=payload.target_amount, current_amount=payload.current_amount, due_date=payload.due_date, priority=payload.priority, status=payload.status)
    db.add(goal)
    db.commit()
    return _goal_payload(goal)


@app.patch("/api/intelligence/goals/{goal_id}")
def update_goal(goal_id: str, request: Request, payload: GoalPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    goal = db.scalar(select(LifeGoalRecord).where(LifeGoalRecord.id == goal_id, LifeGoalRecord.user_id == user.id))
    if goal is None:
        raise HTTPException(status_code=404, detail="Goal not found")
    goal.title, goal.target_amount, goal.current_amount = payload.title, payload.target_amount, payload.current_amount
    goal.due_date, goal.priority, goal.status = payload.due_date, payload.priority, payload.status
    db.commit()
    return _goal_payload(goal)


@app.delete("/api/intelligence/goals/{goal_id}")
def delete_goal(goal_id: str, request: Request, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    goal = db.scalar(select(LifeGoalRecord).where(LifeGoalRecord.id == goal_id, LifeGoalRecord.user_id == user.id))
    if goal is None:
        raise HTTPException(status_code=404, detail="Goal not found")
    db.delete(goal)
    db.commit()
    return {"deleted": True}


@app.post("/api/intelligence/goals/{goal_id}/contributions", status_code=status.HTTP_201_CREATED)
def add_goal_contribution(goal_id: str, request: Request, payload: GoalContributionPayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    goal = db.scalar(select(LifeGoalRecord).where(LifeGoalRecord.id == goal_id, LifeGoalRecord.user_id == user.id))
    if goal is None:
        raise HTTPException(status_code=404, detail="Goal not found")
    contribution = GoalContributionRecord(user_id=user.id, goal_id=goal.id, amount=payload.amount, contributed_at=payload.contributed_at, note=payload.note)
    goal.current_amount += payload.amount
    if goal.current_amount >= goal.target_amount:
        goal.status = "complete"
    db.add(contribution)
    db.commit()
    return {"contributionId": contribution.id, "goal": _goal_payload(goal)}


@app.get("/api/intelligence/goals/{goal_id}/contributions")
def list_goal_contributions(goal_id: str, request: Request, limit: int = 50, offset: int = 0, db: Session = Depends(get_db)):
    user = current_user(request, db)
    goal = db.scalar(select(LifeGoalRecord.id).where(LifeGoalRecord.id == goal_id, LifeGoalRecord.user_id == user.id))
    if goal is None:
        raise HTTPException(status_code=404, detail="Goal not found")
    rows = db.scalars(select(GoalContributionRecord).where(GoalContributionRecord.goal_id == goal_id, GoalContributionRecord.user_id == user.id).order_by(GoalContributionRecord.contributed_at.desc()).offset(max(0, offset)).limit(min(100, max(1, limit)))).all()
    return {"items": [{"id": row.id, "amount": _money(row.amount), "date": row.contributed_at.isoformat(), "note": row.note} for row in rows]}


@app.get("/api/intelligence/rules/triggers")
def list_rule_triggers(request: Request, limit: int = 50, offset: int = 0, db: Session = Depends(get_db)):
    user = current_user(request, db)
    rows = db.scalars(select(TransactionRecord).where(TransactionRecord.user_id == user.id)).all()
    _evaluate_rules(db, user.id, rows)
    triggers = db.scalars(select(RuleTriggerRecord).where(RuleTriggerRecord.user_id == user.id).order_by(RuleTriggerRecord.created_at.desc()).offset(max(0, offset)).limit(min(100, max(1, limit)))).all()
    return {"items": [_trigger_payload(trigger) for trigger in triggers]}


@app.post("/api/intelligence/rules", status_code=status.HTTP_201_CREATED)
def create_rule(request: Request, payload: RulePayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    if payload.condition_type == "category_spend" and payload.category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Choose a valid category for this rule")
    if payload.category is not None and payload.category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Unknown spending category")
    rule = FinancialRuleRecord(user_id=user.id, name=payload.name, condition_type=payload.condition_type, category=payload.category, threshold=payload.threshold, period=payload.period, action=payload.action, active=payload.active)
    db.add(rule)
    db.commit()
    rows = db.scalars(select(TransactionRecord).where(TransactionRecord.user_id == user.id)).all()
    _evaluate_rules(db, user.id, rows)
    return _rule_payload(rule)


@app.patch("/api/intelligence/rules/{rule_id}")
def update_rule(rule_id: str, request: Request, payload: RulePayload, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    if payload.condition_type == "category_spend" and payload.category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Choose a valid category for this rule")
    rule = db.scalar(select(FinancialRuleRecord).where(FinancialRuleRecord.id == rule_id, FinancialRuleRecord.user_id == user.id))
    if rule is None:
        raise HTTPException(status_code=404, detail="Rule not found")
    rule.name, rule.condition_type, rule.category = payload.name, payload.condition_type, payload.category
    rule.threshold, rule.period, rule.action, rule.active = payload.threshold, payload.period, payload.action, payload.active
    db.commit()
    rows = db.scalars(select(TransactionRecord).where(TransactionRecord.user_id == user.id)).all()
    _evaluate_rules(db, user.id, rows)
    return _rule_payload(rule)


@app.delete("/api/intelligence/rules/{rule_id}")
def delete_rule(rule_id: str, request: Request, db: Session = Depends(get_db)):
    user = csrf_user(request, db)
    rule = db.scalar(select(FinancialRuleRecord).where(FinancialRuleRecord.id == rule_id, FinancialRuleRecord.user_id == user.id))
    if rule is None:
        raise HTTPException(status_code=404, detail="Rule not found")
    db.delete(rule)
    db.commit()
    return {"deleted": True}


@app.post("/api/budgets", status_code=status.HTTP_201_CREATED)
def save_budget(request: Request, payload: BudgetPayload, db: Session = Depends(get_db)):
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    user = db.get(User, session.user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    if payload.category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Unknown spending category")
    budget = db.scalar(select(BudgetRecord).where(BudgetRecord.user_id == user.id, BudgetRecord.category == payload.category))
    if budget is None:
        budget = BudgetRecord(user_id=user.id, category=payload.category, limit=payload.limit)
        db.add(budget)
    else:
        budget.limit = payload.limit
    db.commit()
    return {"category": budget.category, "limit": float(budget.limit)}


@app.patch("/api/transactions/{transaction_id}/category")
def update_transaction_category(transaction_id: str, request: Request, payload: CategoryPayload, db: Session = Depends(get_db)):
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    if payload.category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Unknown spending category")
    row = db.scalar(select(TransactionRecord).where(TransactionRecord.id == transaction_id, TransactionRecord.user_id == session.user_id))
    if row is None:
        raise HTTPException(status_code=404, detail="Transaction not found")
    row.category = payload.category
    db.commit()
    return {"id": row.id, "category": row.category}


@app.delete("/api/budgets/{category}")
def remove_budget(category: str, request: Request, db: Session = Depends(get_db)):
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    db.query(BudgetRecord).filter(BudgetRecord.user_id == session.user_id, BudgetRecord.category == category).delete()
    db.commit()
    return {"deleted": True}


@app.post("/api/import/csv", status_code=status.HTTP_201_CREATED)
@limiter.limit("10/hour")
async def import_csv(request: Request, file: Annotated[UploadFile, File()], db: Session = Depends(get_db)):
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    user = db.get(User, session.user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    if not file.filename or not file.filename.lower().endswith((".csv", ".xlsx", ".xls")):
        raise HTTPException(status_code=415, detail="Supported files are CSV, XLSX, and XLS")
    try:
        imported = await run_in_threadpool(import_file_rows, user.id, file.file, file.filename)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    finally:
        await file.close()
    return {"imported": imported}


@app.delete("/api/transactions")
def delete_transactions(request: Request, db: Session = Depends(get_db)):
    session = check_csrf(request, db)
    if session is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    user = db.get(User, session.user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    db.query(TransactionRecord).filter(TransactionRecord.user_id == user.id).delete()
    db.commit()
    return {"deleted": True}
